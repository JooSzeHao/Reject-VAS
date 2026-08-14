"""Geometry + text-fit audit, standing in for a visual render (LibreOffice is unusable here)."""
from pptx import Presentation
from pptx.util import Emu

DECK = '/home/user/Reject-VAS/output/Analisis_Penolakan_VAS_FKP.pptx'
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5

# average glyph width as a fraction of point size, measured conservatively (wide side)
WIDTH_FACTOR = {'Calibri': 0.475, 'Cambria': 0.50}
BOLD_EXTRA = 1.045

prs = Presentation(DECK)
inches = lambda v: Emu(v).inches

problems = []
for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if sh.left is None:
            continue
        x, y = inches(sh.left), inches(sh.top)
        w, h = inches(sh.width), inches(sh.height)
        name = sh.shape_type
        if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
            problems.append(f'S{si}: OUT OF BOUNDS {name} at ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}')
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if not txt:
            continue
        boxes.append((x, y, w, h, txt[:34]))
        # text fit: longest paragraph wrapped into the box width
        for para in sh.text_frame.paragraphs:
            runs = [r for r in para.runs if r.text]
            if not runs:
                continue
            size = max((r.font.size.pt if r.font.size else 18) for r in runs)
            face = runs[0].font.name or 'Calibri'
            bold = any(r.font.bold for r in runs)
            fac = WIDTH_FACTOR.get(face, 0.48) * (BOLD_EXTRA if bold else 1.0)
            text = ''.join(r.text for r in runs)
            char_w = size * fac / 72.0          # inches per character
            usable = w - 0.2                     # pptx internal padding both sides
            per_line = max(1, int(usable / char_w))
            lines = 0
            for hard in text.split('\n'):
                words, cur = hard.split(' '), ''
                ln = 1
                for word in words:
                    trial = (cur + ' ' + word).strip()
                    if len(trial) > per_line and cur:
                        ln += 1
                        cur = word
                    else:
                        cur = trial
                lines += ln
            needed = lines * size * 1.28 / 72.0
            if needed > h + 0.02:
                problems.append(
                    f'S{si}: TEXT OVERFLOW ~{needed:.2f}in in {h:.2f}in box '
                    f'({size:.0f}pt, {lines} lines) — "{text[:52]}"')

    # overlap between text boxes
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
            oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
            if ox > 0.05 and oy > 0.05:
                problems.append(f'S{si}: TEXT OVERLAP {ox:.2f}x{oy:.2f}in — "{a[4]}" / "{b[4]}"')

print(f'{len(prs.slides.__iter__.__self__._sldIdLst)} slides checked')
print('\n'.join(problems) if problems else 'No geometry or text-fit problems found.')
